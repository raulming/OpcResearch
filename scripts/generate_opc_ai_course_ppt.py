from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


OUT = Path("docs/OPC-AI公益课-自测评估与成长路径.pptx")

WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

BG = RGBColor(7, 16, 30)
PANEL = RGBColor(13, 30, 50)
PANEL_2 = RGBColor(18, 43, 68)
CYAN = RGBColor(39, 224, 255)
GREEN = RGBColor(73, 236, 169)
AMBER = RGBColor(255, 190, 84)
RED = RGBColor(255, 91, 110)
VIOLET = RGBColor(154, 128, 255)
WHITE = RGBColor(238, 247, 255)
MUTED = RGBColor(154, 178, 205)
GRID = RGBColor(24, 58, 86)

FONT = "Microsoft YaHei"


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color=CYAN, width=1.2, transparency=0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_bg(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    # Subtle tech grid
    for i in range(0, 14):
        x = Inches(i)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, 0, Pt(0.5), WIDE_H)
        set_fill(line, GRID, 65)
        line.line.fill.background()
    for j in range(0, 8):
        y = Inches(j)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, y, WIDE_W, Pt(0.5))
        set_fill(line, GRID, 70)
        line.line.fill.background()

    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.3), Inches(-1.1), Inches(4.4), Inches(4.4))
    set_fill(glow, RGBColor(0, 144, 190), 78)
    glow.line.fill.background()

    glow2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.5), Inches(5.6), Inches(4.6), Inches(2.4))
    set_fill(glow2, RGBColor(52, 206, 159), 86)
    glow2.line.fill.background()


def add_textbox(slide, x, y, w, h, text, size=22, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.65), Inches(0.42), Inches(9.8), Inches(0.52), title, 26, WHITE, True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.02), Inches(1.25), Pt(3))
    set_fill(accent, CYAN)
    accent.line.fill.background()
    if subtitle:
        add_textbox(slide, Inches(0.65), Inches(1.12), Inches(10.8), Inches(0.35), subtitle, 11, MUTED)


def add_footer(slide, idx):
    add_textbox(slide, Inches(11.35), Inches(7.04), Inches(1.2), Inches(0.22), f"{idx:02d}", 9, MUTED, False, PP_ALIGN.RIGHT)
    mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.72), Inches(0.55), Pt(3), Inches(6.25))
    set_fill(mark, CYAN, 15)
    mark.line.fill.background()


def add_card(slide, x, y, w, h, title, body, color=CYAN):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    set_fill(card, PANEL, 8)
    set_line(card, color, 1.0, 12)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(3))
    set_fill(bar, color, 0)
    bar.line.fill.background()
    add_textbox(slide, x + Inches(0.22), y + Inches(0.18), w - Inches(0.44), Inches(0.34), title, 15, WHITE, True)
    add_textbox(slide, x + Inches(0.22), y + Inches(0.62), w - Inches(0.44), h - Inches(0.78), body, 10.5, MUTED)


def add_bullets(slide, x, y, w, h, bullets, size=16, color=WHITE):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return box


def add_table(slide, x, y, w, h, rows, cols, data, col_widths=None, header=True):
    shape = slide.shapes.add_table(rows, cols, x, y, w, h)
    table = shape.table
    if col_widths:
        for i, width in enumerate(col_widths):
            table.columns[i].width = width
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            fill = PANEL_2 if r == 0 and header else PANEL
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.text_frame.word_wrap = True
            cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(9.5 if r else 10.5)
                p.font.bold = bool(r == 0 and header)
                p.font.color.rgb = WHITE if r == 0 and header else MUTED
    return shape


def add_score_band(slide, x, y, w, h, label, score, suggestion, color):
    band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    set_fill(band, PANEL, 7)
    set_line(band, color, 1.0, 18)
    add_textbox(slide, x + Inches(0.18), y + Inches(0.13), Inches(1.25), Inches(0.3), score, 14, color, True)
    add_textbox(slide, x + Inches(1.45), y + Inches(0.13), Inches(1.45), Inches(0.3), label, 13, WHITE, True)
    add_textbox(slide, x + Inches(3.05), y + Inches(0.13), w - Inches(3.25), Inches(0.3), suggestion, 10.5, MUTED)


def add_compact_checklist(slide, x, y, items, color):
    for i, item in enumerate(items):
        row_y = y + Inches(i * 0.47)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, row_y + Inches(0.06), Inches(0.16), Inches(0.16))
        set_fill(dot, color, 0)
        dot.line.fill.background()
        add_textbox(slide, x + Inches(0.28), row_y, Inches(4.8), Inches(0.24), item, 10.2, WHITE)


def persona_slide(prs, idx, name, who, need, pain, weak, scenarios, advice, color):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, name, who)
    add_footer(slide, idx)
    add_card(slide, Inches(0.75), Inches(1.65), Inches(3.7), Inches(1.45), "核心诉求", need, color)
    add_card(slide, Inches(4.75), Inches(1.65), Inches(3.7), Inches(1.45), "典型问题", pain, AMBER)
    add_card(slide, Inches(8.75), Inches(1.65), Inches(3.7), Inches(1.45), "能力短板", weak, RED)
    add_textbox(slide, Inches(0.75), Inches(3.42), Inches(2.8), Inches(0.32), "适合切入场景", 16, WHITE, True)
    add_bullets(slide, Inches(0.9), Inches(3.9), Inches(5.2), Inches(2.25), scenarios, 13, WHITE)
    quote = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.45), Inches(3.65), Inches(5.95), Inches(2.3))
    set_fill(quote, RGBColor(9, 27, 45), 0)
    set_line(quote, color, 1.2, 5)
    add_textbox(slide, Inches(6.82), Inches(4.05), Inches(5.15), Inches(1.3), advice, 18, WHITE, True)


def main():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H

    # 1 Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_textbox(slide, Inches(0.75), Inches(0.72), Inches(3.7), Inches(0.36), "公益课 / AI 能力评估", 13, CYAN, True)
    add_textbox(slide, Inches(0.75), Inches(1.55), Inches(10.8), Inches(1.65), "OPC 创业自测\n与 AI 能力成长路径", 38, WHITE, True)
    add_textbox(slide, Inches(0.82), Inches(3.55), Inches(8.8), Inches(0.55), "从一时冲动，到可执行的个人业务闭环", 20, GREEN, True)
    add_card(slide, Inches(0.82), Inches(5.18), Inches(3.55), Inches(1.0), "课程目标", "识别人群差异，评估 AI 能力，找到每个人的第一步行动。", CYAN)
    add_card(slide, Inches(4.62), Inches(5.18), Inches(3.55), Inches(1.0), "适合对象", "职场人、创业者、创作者、管理者、技术人、学习者。", GREEN)
    add_card(slide, Inches(8.42), Inches(5.18), Inches(3.55), Inches(1.0), "关键词", "自测 / 风险 / 场景 / 交付 / 产品化", AMBER)
    add_footer(slide, 1)

    # 2 Agenda
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "今天要解决的 4 个问题")
    add_footer(slide, 2)
    items = [
        ("01", "AI 对 OPC 的真正价值", "不是替代人，而是放大一个人的业务闭环能力。"),
        ("02", "AI 能力如何评估", "用 6 个维度判断当前能力位置。"),
        ("03", "不同人群如何画像", "根据需求、短板、场景设计差异化路径。"),
        ("04", "下一步如何行动", "从一个具体场景开始，形成可复用工作流。"),
    ]
    for i, (num, title, body) in enumerate(items):
        x = Inches(0.8 + i * 3.05)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.65), Inches(3.2))
        set_fill(card, PANEL, 4)
        set_line(card, [CYAN, GREEN, AMBER, VIOLET][i], 1.3, 0)
        add_textbox(slide, x + Inches(0.22), Inches(2.22), Inches(0.7), Inches(0.35), num, 18, [CYAN, GREEN, AMBER, VIOLET][i], True)
        add_textbox(slide, x + Inches(0.22), Inches(2.86), Inches(2.15), Inches(0.62), title, 18, WHITE, True)
        add_textbox(slide, x + Inches(0.22), Inches(3.75), Inches(2.15), Inches(1.05), body, 12, MUTED)

    # 3 Core insight
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "核心判断：AI 放大的是业务闭环")
    add_footer(slide, 3)
    add_textbox(slide, Inches(1.05), Inches(1.75), Inches(11.1), Inches(0.75), "适合做 OPC 的人，不一定最懂技术，但一定能用 AI 完成一个完整闭环。", 26, WHITE, True, PP_ALIGN.CENTER)
    chain = ["发现问题", "设计方案", "生产内容", "交付服务", "持续迭代"]
    colors = [CYAN, GREEN, AMBER, VIOLET, RED]
    for i, label in enumerate(chain):
        x = Inches(0.9 + i * 2.45)
        node = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, x, Inches(3.45), Inches(1.8), Inches(1.1))
        set_fill(node, PANEL_2, 0)
        set_line(node, colors[i], 1.5, 0)
        add_textbox(slide, x + Inches(0.17), Inches(3.83), Inches(1.46), Inches(0.28), label, 14, WHITE, True, PP_ALIGN.CENTER)
        if i < 4:
            add_textbox(slide, x + Inches(1.85), Inches(3.83), Inches(0.42), Inches(0.28), "→", 22, MUTED, True, PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.45), Inches(5.35), Inches(10.4), Inches(0.65), "评估 AI 能力，本质上是在评估：这个人能否借助 AI，把想法变成稳定产出。", 18, GREEN, True, PP_ALIGN.CENTER)

    # 4 Model
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "AI 能力评估六维模型")
    add_footer(slide, 4)
    dims = [
        ("问题定义", "说清目标、对象、限制、结果标准", CYAN),
        ("信息判断", "识别幻觉、查证来源、对比观点", GREEN),
        ("提问协作", "拆任务、给背景、要格式、会迭代", AMBER),
        ("业务转化", "把 AI 结果变成交付物或服务", VIOLET),
        ("工具整合", "文本、表格、图片、PPT、自动化组合", CYAN),
        ("执行复盘", "形成模板、记录反馈、持续优化", GREEN),
    ]
    for i, (t, b, c) in enumerate(dims):
        x = Inches(0.75 + (i % 3) * 4.15)
        y = Inches(1.7 + (i // 3) * 2.15)
        add_card(slide, x, y, Inches(3.65), Inches(1.55), t, b, c)

    # 5 Score table
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "现场自测：10 道题，每题 1-5 分")
    add_footer(slide, 5)
    questions = [
        "我能清楚描述自己想解决的问题",
        "我能给 AI 提供背景、目标、限制条件",
        "我能判断 AI 输出中哪些内容不可靠",
        "我能让 AI 生成可直接使用的文档/表格/PPT",
        "我能把 AI 结果修改成适合自己业务的版本",
        "我至少掌握 3 个以上 AI 应用场景",
        "我能把重复工作做成模板或流程",
        "我能用 AI 辅助内容、客户沟通或产品设计",
        "我能持续复盘 AI 产出的效果",
        "我知道下一步要改造哪个具体工作环节",
    ]
    for i, q in enumerate(questions):
        x = Inches(0.88 if i < 5 else 6.78)
        y = Inches(1.72 + (i % 5) * 0.82)
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.35), Inches(0.5))
        set_fill(pill, PANEL, 8)
        set_line(pill, CYAN if i < 5 else GREEN, 0.8, 40)
        add_textbox(slide, x + Inches(0.16), y + Inches(0.12), Inches(0.38), Inches(0.2), str(i + 1), 10, CYAN, True)
        add_textbox(slide, x + Inches(0.58), y + Inches(0.1), Inches(4.55), Inches(0.28), q, 10.5, WHITE)
    add_textbox(slide, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.38), "评分方式：1 分完全不会，3 分偶尔能做到，5 分可以稳定做到。", 15, AMBER, True, PP_ALIGN.CENTER)

    # 6 Score levels
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "分数解释：四类 AI 使用者")
    add_footer(slide, 6)
    data = [
        ["总分", "类型", "典型状态", "建议路径"],
        ["10-20", "入门型", "知道 AI 重要，但还没有稳定使用习惯", "从日常提效开始"],
        ["21-35", "应用型", "能完成一些任务，但流程不稳定", "建立个人 AI 工作流"],
        ["36-45", "生产型", "能产出内容、方案或服务交付物", "尝试产品化与标准化"],
        ["46-50", "操盘型", "能用 AI 组织业务闭环和自动化系统", "搭建 OPC 业务系统"],
    ]
    add_table(slide, Inches(0.78), Inches(1.75), Inches(11.8), Inches(3.6), 5, 4, data, [Inches(1.35), Inches(1.55), Inches(4.2), Inches(4.7)])
    add_textbox(slide, Inches(1.05), Inches(5.85), Inches(11.1), Inches(0.55), "重点不是比较谁分高，而是找到每个人最值得开始的 AI 场景。", 20, GREEN, True, PP_ALIGN.CENTER)

    # 7 Survey report: founder readiness
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "调研报告页 1：OPC 启动准备度测试", "现场填写后，把结果分成“冲动探索、副业验证、服务启动、系统放大”四种状态。")
    add_footer(slide, 7)
    readiness = [
        ["维度", "自测问题", "收集信息"],
        ["动机", "是长期问题驱动，还是短期焦虑驱动？", "创业动机 / 时间窗口"],
        ["客户", "我能清楚说出目标客户是谁吗？", "目标客户 / 人群标签"],
        ["痛点", "客户是否有高频、刚需、有成本的问题？", "具体痛点 / 当前代价"],
        ["反馈", "我是否和真实潜在客户聊过？", "访谈记录 / 明确反馈"],
        ["付费", "客户为什么愿意付费，大概能付多少？", "付费理由 / 价格带"],
        ["获客", "30 天内能触达 50-100 个潜在客户吗？", "渠道 / 私域 / 触达方式"],
        ["交付", "能否做成清晰服务包或产品包？", "服务边界 / 报价 / SOP"],
        ["风险", "是否有 3-6 个月验证期和止损线？", "现金流 / 止损线"],
    ]
    add_table(slide, Inches(0.65), Inches(1.48), Inches(7.25), Inches(5.15), 9, 3, readiness, [Inches(1.05), Inches(3.9), Inches(2.3)])
    add_textbox(slide, Inches(8.25), Inches(1.55), Inches(3.9), Inches(0.36), "评分解释", 17, WHITE, True)
    add_score_band(slide, Inches(8.25), Inches(2.05), Inches(4.15), Inches(0.56), "冲动探索", "8-16", "先做职业提效和客户访谈", RED)
    add_score_band(slide, Inches(8.25), Inches(2.78), Inches(4.15), Inches(0.56), "副业验证", "17-24", "用 30 天做最小验证", AMBER)
    add_score_band(slide, Inches(8.25), Inches(3.51), Inches(4.15), Inches(0.56), "服务启动", "25-30", "设计服务包并找真实试单", GREEN)
    add_score_band(slide, Inches(8.25), Inches(4.24), Inches(4.15), Inches(0.56), "系统放大", "31-40", "沉淀 SOP、产品和渠道系统", CYAN)
    add_card(slide, Inches(8.25), Inches(5.18), Inches(4.15), Inches(1.04), "课堂动作", "每题 1-5 分。重点不是证明自己适合创业，而是找到下一步最小验证。", VIOLET)

    # 8 Survey report: AI capability and startup direction
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "调研报告页 2：我的 AI 状态适合从哪里启动？", "把 AI 能力分数和启动准备度合并，给出第一场景和后续服务线索。")
    add_footer(slide, 8)
    add_textbox(slide, Inches(0.75), Inches(1.48), Inches(3.9), Inches(0.34), "先看两个分数", 17, WHITE, True)
    add_card(slide, Inches(0.75), Inches(1.98), Inches(2.95), Inches(1.18), "启动准备度", "8 个维度，每题 1-5 分。\n判断现在是探索、验证、启动还是放大。", AMBER)
    add_card(slide, Inches(4.0), Inches(1.98), Inches(2.95), Inches(1.18), "AI 能力分", "10 道题，每题 1-5 分。\n判断是入门、应用、生产还是操盘。", CYAN)
    direction_data = [
        ["组合状态", "最适合的启动方向", "第一份交付物", "后续服务线索"],
        ["准备低 + AI 低", "职业提效 / 客户访谈", "工作流清单 + 10 个访谈问题", "入门训练营"],
        ["准备中 + AI 中", "副业验证 / 内容获客", "30 天验证计划 + 1 个试单方案", "场景工作坊"],
        ["准备高 + AI 中", "服务型 OPC", "服务包 + 报价单 + 交付 SOP", "咨询陪跑"],
        ["准备中高 + AI 高", "工具化 / 产品化", "模板产品 / 小工具 / 付费测试页", "产品共创"],
        ["准备高 + AI 高", "一人公司系统", "获客-交付-复盘自动化流程", "系统搭建"],
    ]
    add_table(slide, Inches(0.75), Inches(3.55), Inches(11.8), Inches(2.75), 6, 4, direction_data, [Inches(2.05), Inches(2.85), Inches(3.3), Inches(3.6)])
    add_textbox(slide, Inches(7.35), Inches(1.48), Inches(3.9), Inches(0.34), "建议收集 6 类信息", 17, WHITE, True)
    add_compact_checklist(
        slide,
        Inches(7.42),
        Inches(1.98),
        ["当前职业 / 专业能力", "想服务的人群与痛点", "每周可投入时间", "可接受投入预算", "已有渠道或私域规模", "最想尝试的 AI 场景"],
        GREEN,
    )

    # 9 Persona overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "六类人群画像总览")
    add_footer(slide, 9)
    personas = [
        ("职场转型型", "保饭碗 / 提效率 / 找副业"),
        ("个体创业型", "获客 / 交付 / 标准化"),
        ("老板管理型", "降本增效 / 团队改造"),
        ("内容创作者型", "持续产出 / 定位 / 转化"),
        ("技术产品型", "工具化 / 自动化 / 商业验证"),
        ("学习成长型", "入门 / 学习 / 找方向"),
    ]
    for i, (t, b) in enumerate(personas):
        x = Inches(0.85 + (i % 3) * 4.15)
        y = Inches(1.7 + (i // 3) * 2.1)
        add_card(slide, x, y, Inches(3.55), Inches(1.45), t, b, [CYAN, GREEN, AMBER, VIOLET, CYAN, GREEN][i])

    # Personas 10-15
    persona_slide(
        prs, 10, "画像一：职场转型型", "行政、人事、运营、销售、客服、传统行业从业者",
        "担心被 AI 淘汰，希望提升竞争力，发展副业。",
        "会用一点 AI，但常停留在写文案、写总结。",
        "问题定义弱；工具整合弱；业务转化意识不足。",
        ["简历优化", "工作汇报", "SOP 编写", "客户话术", "资料整理", "副业内容生产"],
        "先把 AI 用进每天的工作流，从“会问 AI”变成“用 AI 交付更好的工作成果”。",
        CYAN,
    )
    persona_slide(
        prs, 11, "画像二：个体创业 / 自由职业型", "咨询师、设计师、培训师、摄影师、教练、自由顾问",
        "提升获客、内容生产和客户交付效率。",
        "专业能力有，但表达、包装、营销和交付标准化不足。",
        "产品化不足；流程沉淀不足；内容矩阵不足。",
        ["客户画像分析", "内容选题库", "咨询方案", "报告模板", "课程大纲", "服务 SOP"],
        "让 AI 帮你把专业表达清楚、包装清楚、交付稳定，而不是替代你的专业。",
        GREEN,
    )
    persona_slide(
        prs, 12, "画像三：老板 / 管理者型", "中小企业老板、门店老板、团队负责人、项目负责人",
        "降本增效，让团队真正把 AI 用起来。",
        "自己知道 AI 重要，但团队不会用，也缺落地机制。",
        "场景识别弱；流程改造弱；团队规范不足。",
        ["销售话术库", "客户 FAQ", "员工培训资料", "会议纪要", "竞品分析", "经营数据解读"],
        "老板学 AI，是为了重新设计团队工作方式，优先改造高频、重复、标准化环节。",
        AMBER,
    )
    persona_slide(
        prs, 13, "画像四：内容创作者型", "自媒体人、博主、主播、知识付费从业者、短视频创作者",
        "持续产出内容，提升传播和商业转化。",
        "灵感不稳定，内容同质化，有流量但产品承接弱。",
        "选题判断弱；用户洞察弱；商业闭环不清。",
        ["选题库", "标题分析", "脚本生成", "图文改写", "直播大纲", "私域转化话术"],
        "AI 能放大产量，但真正决定价值的是定位、观点和用户理解。",
        VIOLET,
    )
    persona_slide(
        prs, 14, "画像五：技术 / 产品型", "程序员、产品经理、数据分析师、自动化爱好者",
        "用 AI 做工具、产品、自动化系统。",
        "技术强，但容易先做工具，后找需求。",
        "用户洞察不足；商业化能力不足；表达销售较弱。",
        ["AI Agent", "自动化报表", "内部工具", "数据分析", "行业小工具", "SaaS 原型验证"],
        "不要先问我能做什么工具，而要先问谁愿意为什么问题付钱。",
        CYAN,
    )
    persona_slide(
        prs, 15, "画像六：学习成长型", "学生、宝妈、退休人群、转行探索者、AI 初学者",
        "了解 AI，提升学习效率，找到个人可能性。",
        "害怕技术门槛，不知道从哪里开始，容易被工具信息淹没。",
        "基础认知不足；缺少稳定场景；容易浅尝辄止。",
        ["学习计划", "读书总结", "英语学习", "资料整理", "生活规划", "兴趣内容创作"],
        "第一阶段不要焦虑变现，先建立 AI 使用习惯，让 AI 成为学习和表达伙伴。",
        GREEN,
    )

    # 16 Scenario map
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "不同人群的 AI 第一场景")
    add_footer(slide, 16)
    data = [
        ["人群", "最优先场景", "第一份交付物"],
        ["职场转型型", "日常工作提效", "一份工作汇报 / SOP / 简历优化稿"],
        ["个体创业型", "获客与交付标准化", "一套服务介绍 + 咨询方案模板"],
        ["老板管理型", "团队高频重复流程", "销售话术库 / FAQ / 会议行动清单"],
        ["内容创作者型", "选题与内容生产", "30 天选题库 + 3 条脚本"],
        ["技术产品型", "需求验证与原型", "一个可演示的小工具或自动化流程"],
        ["学习成长型", "学习与表达", "个人学习计划 + 读书总结模板"],
    ]
    add_table(slide, Inches(0.85), Inches(1.55), Inches(11.65), Inches(4.65), 7, 3, data, [Inches(2.0), Inches(3.65), Inches(6.0)])

    # 17 Growth path
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "OPC 的 AI 成长四阶段")
    add_footer(slide, 17)
    stages = [
        ("会用 AI", "基础提问、总结整理、内容生成", "效率提升"),
        ("用 AI 做交付", "方案、报告、客户沟通、内容生产", "稳定产出"),
        ("用 AI 建系统", "获客、交付、复盘、自动化", "业务闭环"),
        ("用 AI 产品化", "模板化、课程化、工具化、社群化", "降低边际成本"),
    ]
    for i, (t, b, r) in enumerate(stages):
        x = Inches(0.85 + i * 3.08)
        y = Inches(2.05 + i * 0.35)
        add_card(slide, x, y, Inches(2.68), Inches(2.0), f"{i + 1}. {t}", f"{b}\n\n结果：{r}", [CYAN, GREEN, AMBER, VIOLET][i])
    add_textbox(slide, Inches(1.05), Inches(5.82), Inches(11.1), Inches(0.5), "从“卖时间”走向“卖产品、卖系统、卖可复制能力”。", 20, GREEN, True, PP_ALIGN.CENTER)

    # 18 Closing
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_footer(slide, 18)
    add_textbox(slide, Inches(0.9), Inches(1.05), Inches(11.55), Inches(1.65), "不要问：AI 能替我做什么？", 36, WHITE, True, PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.1), Inches(3.0), Inches(11.1), Inches(1.25), "要问：我想创造什么价值，\nAI 能帮我放大哪一部分？", 32, GREEN, True, PP_ALIGN.CENTER)
    add_card(slide, Inches(2.0), Inches(5.35), Inches(9.35), Inches(0.9), "课后行动", "选一个最真实、最高频、最困扰你的场景，用 AI 做出第一份可交付成果。", CYAN)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
